#!/usr/bin/env python3
"""Build fact-bounded CV and optional letter in DOCX and PDF from JSON."""
import argparse
import hashlib
import json
import os
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader
from xml.sax.saxutils import escape


SECTION_ALIASES = {
    'PROFILE': 'PROFILE',
    'SUMMARY': 'PROFILE',
    'PROFESSIONAL SUMMARY': 'PROFILE',
    'SELECTED AI & AGENTIC SYSTEMS WORK': 'SELECTED AI & AGENTIC SYSTEMS WORK',
    'CORE EXPERTISE': 'CORE EXPERTISE',
    'CORE SKILLS': 'CORE EXPERTISE',
    'AREAS OF EXPERTISE': 'CORE EXPERTISE',
    'TECHNOLOGY': 'TECHNOLOGY',
    'TECHNICAL SKILLS': 'TECHNOLOGY',
    'TECH STACK': 'TECHNOLOGY',
    'PROFESSIONAL EXPERIENCE': 'PROFESSIONAL EXPERIENCE',
    'WORK EXPERIENCE': 'PROFESSIONAL EXPERIENCE',
    'EXPERIENCE': 'PROFESSIONAL EXPERIENCE',
    'EMPLOYMENT HISTORY': 'PROFESSIONAL EXPERIENCE',
    'EARLIER EXPERIENCE': 'EARLIER EXPERIENCE',
    'ADDITIONAL EXPERIENCE': 'EARLIER EXPERIENCE',
    'EDUCATION': 'EDUCATION',
    'ACADEMIC BACKGROUND': 'EDUCATION',
    'SKILLS': 'SKILLS',
    'CONTACT': 'CONTACT',
    'CONTACT DETAILS': 'CONTACT',
}
DATE_LINE = re.compile(
    r'^(?:(?:[A-Za-z]{3,9}\s+)?\d{4})(?:\s*[–—-]\s*(?:(?:[A-Za-z]{3,9}\s+)?(?:\d{4}|present|current)))?$|^\d{1,2}/\d{4}\s*[–—-]',
    re.IGNORECASE,
)


def load(path):
    value = json.loads(Path(path).read_text(encoding='utf-8'))
    if not isinstance(value, dict):
        raise ValueError(f'{path}: expected object')
    return value


def required(obj, keys, label):
    for key in keys:
        if not isinstance(obj.get(key), str) or not obj[key].strip():
            raise ValueError(f'{label}: missing {key}')


def _source_rows(path):
    """Return normalized text rows and locators for supported CV formats."""
    path = Path(path)
    suffix = path.suffix.casefold()
    if suffix == '.docx':
        doc = Document(path)
        rows = []
        for i, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name
            if style.startswith('List Bullet') or style.startswith('List Number'):
                text = '• ' + text
            rows.append((i, text, style, {'kind': 'paragraph', 'index': i}))
        return rows, 'docx'

    if suffix in {'.md', '.markdown', '.txt'}:
        text = path.read_text(encoding='utf-8-sig')
        lines = text.splitlines()
        kind = 'markdown' if suffix in {'.md', '.markdown'} else 'text'
        return _text_rows(lines, kind)

    if suffix == '.pdf':
        pages = PdfReader(str(path)).pages
        lines = []
        for page_number, page in enumerate(pages, 1):
            page_text = page.extract_text() or ''
            if not page_text.strip():
                continue
            for line_number, line in enumerate(page_text.splitlines(), 1):
                lines.append((line, {'kind': 'page_line', 'page': page_number, 'line': line_number}))
        if not any(line.strip() for line, _ in lines):
            raise ValueError('PDF has no selectable text; provide DOCX/Markdown or OCR and review it first')
        return _text_rows(lines, 'pdf')

    raise ValueError('supported CV formats are DOCX, Markdown (.md), text (.txt), and selectable-text PDF')


def _clean_markdown(text):
    text = re.sub(r'^\s{0,3}#{1,6}\s+', '', text)
    text = re.sub(r'!\[([^\]]*)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    text = re.sub(r'(?<!\\)[*_`~]', '', text)
    return text.strip()


def _section_name(text):
    normalized = re.sub(r'\s+', ' ', text.strip().upper())
    return SECTION_ALIASES.get(normalized)


def _text_rows(lines, kind):
    rows = []
    current_section = None
    first_content = True
    for index, item in enumerate(lines):
        if kind == 'pdf':
            raw, locator = item
        else:
            raw = item
            locator = {'kind': 'line', 'index': index + 1}
        raw = raw.strip()
        if not raw or re.fullmatch(r'[-*_\s]{3,}', raw):
            continue
        heading_level = len(raw) - len(raw.lstrip('#')) if kind == 'markdown' else 0
        text = _clean_markdown(raw) if kind == 'markdown' else raw
        section = _section_name(text)
        if first_content and (kind == 'pdf' or heading_level == 1) and not section:
            style = 'Title'
        elif section:
            text = section
            style = 'Heading 1'
            current_section = section
        elif heading_level >= 2:
            style = 'Heading 2' if heading_level >= 3 else 'Heading 1'
            if style == 'Heading 1':
                current_section = None
        else:
            style = 'Normal'
            if re.match(r'^\s*(?:[-*+] |\d+[.)] )', text):
                text = '• ' + re.sub(r'^\s*(?:[-*+] |\d+[.)] )', '', text)
            if current_section == 'PROFESSIONAL EXPERIENCE':
                if DATE_LINE.fullmatch(text):
                    style = 'RoleMeta'
                elif re.search(r'\s(?:—|–|-)\s', text) and len(text) < 180:
                    style = 'Heading 2'
        if kind == 'markdown' and raw.startswith('#'):
            # Preserve Markdown section hierarchy while stripping its marker.
            if section:
                style = 'Heading 1'
            elif heading_level == 1 and first_content:
                style = 'Title'
            elif heading_level >= 2 and not section:
                style = 'Heading 2' if heading_level >= 3 else 'Heading 1'
        rows.append((index, text, style, locator))
        first_content = False
    return rows, kind


def import_source(path):
    """Import a structured DOCX, Markdown, text PDF, or text file; fail closed on unknown structure."""
    path = Path(path).resolve()
    rows, source_format = _source_rows(path)
    sections = {'PROFILE', 'SELECTED AI & AGENTIC SYSTEMS WORK', 'CORE EXPERTISE',
                'TECHNOLOGY', 'PROFESSIONAL EXPERIENCE', 'EARLIER EXPERIENCE', 'EDUCATION', 'SKILLS', 'CONTACT'}
    profile = {'source': {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                          'format': source_format},
               'contact': {}, 'facts': [], 'roles': [], 'highlights': [],
               'core_expertise': [], 'technology': [], 'earlier_experience': [],
               'education': [], 'skills': [], 'provenance': {}}
    section = None
    role = None
    for i, text, style, locator in rows:
        provenance = locator
        if style == 'Title':
            profile['name'] = text
            profile['provenance']['name'] = provenance
        elif style == 'Heading 1':
            section = text.upper()
            if section not in sections:
                raise ValueError(f'unrecognized source section at {locator}: {text}')
            role = None
        elif section is None:
            if 'headline' not in profile:
                profile['headline'] = text
                profile['provenance']['headline'] = provenance
            else:
                # Do not infer an absent email address from a web profile.
                for part in (x.strip() for x in text.split('|')):
                    if '@' in part: key = 'email'
                    elif part.startswith('+'): key = 'phone'
                    elif 'linkedin.com/' in part: key = 'linkedin'
                    elif 'github.com/' in part: key = 'github'
                    else: key = 'location'
                    profile['contact'][key] = part
                    profile['provenance']['contact.' + key] = provenance
        elif section == 'PROFESSIONAL EXPERIENCE':
            if style == 'Heading 2':
                separator = next((marker for marker in (' — ', ' – ', ' - ') if marker in text), None)
                if separator is None:
                    raise ValueError(f'unrecognized role heading at {locator}')
                title, company = text.split(separator, 1)
                role = {'title': title, 'company': company, 'dates': '', 'fact_ids': [],
                        'source_locations': {'heading': provenance}}
                if source_format == 'docx':
                    role['source_paragraphs'] = {'heading': i}
                profile['roles'].append(role)
            elif role is None:
                raise ValueError(f'role content without heading at {locator}')
            elif style == 'RoleMeta' or (not role['dates'] and DATE_LINE.fullmatch(text)):
                role['dates'] = text
                role['source_locations']['dates'] = provenance
                if source_format == 'docx':
                    role['source_paragraphs']['dates'] = i
            elif text.startswith('•'):
                fid = f"r{len(profile['roles'])}_{len(role['fact_ids']) + 1}"
                fact = {'id': fid, 'text': text.lstrip('• ').strip(), 'source_location': provenance}
                if source_format == 'docx':
                    fact['source_paragraph'] = i
                profile['facts'].append(fact)
                role['fact_ids'].append(fid)
            elif 'narrative' not in role:
                role['narrative'] = text
                role['source_locations']['narrative'] = provenance
                if source_format == 'docx':
                    role['source_paragraphs']['narrative'] = i
            else:
                raise ValueError(f'unrecognized role content at {locator}: {text}')
        elif section == 'PROFILE':
            profile['profile_summary'] = '\n'.join(filter(None, (profile.get('profile_summary'), text)))
            profile['provenance'].setdefault('profile_summary', []).append(provenance)
        elif section == 'SELECTED AI & AGENTIC SYSTEMS WORK':
            item = {'id': f"h{len(profile['highlights']) + 1}",
                    'text': text.lstrip('• ').strip(), 'source_location': provenance}
            if source_format == 'docx':
                item['source_paragraph'] = i
            profile['highlights'].append(item)
        elif section == 'CORE EXPERTISE':
            profile['core_expertise'].extend(x.strip() for x in text.split('•') if x.strip())
            profile['provenance'].setdefault('core_expertise', []).append(provenance)
        elif section == 'TECHNOLOGY':
            profile['technology'].append(text)
            profile['provenance'].setdefault('technology', []).append(provenance)
        elif section == 'SKILLS':
            profile['skills'].extend(x.strip() for x in re.split(r'[•|,]', text.lstrip('• ').strip()) if x.strip())
            profile['provenance'].setdefault('skills', []).append(provenance)
        elif section == 'CONTACT':
            for part in (x.strip() for x in text.split('|')):
                if not part:
                    continue
                if '@' in part: key = 'email'
                elif part.startswith('+'): key = 'phone'
                elif 'linkedin.com/' in part: key = 'linkedin'
                elif 'github.com/' in part: key = 'github'
                else: key = 'location'
                profile['contact'][key] = part
                profile['provenance']['contact.' + key] = provenance
        elif section == 'EARLIER EXPERIENCE':
            profile['earlier_experience'].append(text)
            profile['provenance'].setdefault('earlier_experience', []).append(provenance)
        elif section == 'EDUCATION':
            profile['education'].append(text)
            profile['provenance'].setdefault('education', []).append(provenance)
    if not profile['roles'] or any(not r['dates'] or not r['fact_ids'] for r in profile['roles']):
        raise ValueError('source roles incomplete')
    return profile


def assemble(profile, job, selection):
    required(profile, ['name', 'headline'], 'profile')
    required(job, ['id', 'title', 'company', 'url', 'location', 'work_model'], 'job')
    if not job['url'].startswith(('https://', 'http://')):
        raise ValueError('job: URL must be absolute')
    contact = profile.get('contact', {})
    required(contact, ['phone', 'location'], 'contact')
    facts = profile.get('facts')
    roles = profile.get('roles')
    if not isinstance(facts, list) or not isinstance(roles, list) or not roles:
        raise ValueError('profile: nonempty roles and facts lists required')
    by_id = {}
    for fact in facts:
        required(fact, ['id', 'text'], 'fact')
        if fact['id'] in by_id:
            raise ValueError('duplicate fact id: ' + fact['id'])
        by_id[fact['id']] = fact['text']
    requested = selection.get('role_fact_ids', {})
    if not isinstance(requested, dict) or any(not str(i).isdigit() or int(i) >= len(roles) for i in requested):
        raise ValueError('selection: invalid role index')
    rendered_roles = []
    for i, role in enumerate(roles):
        required(role, ['title', 'company', 'dates'], f'role {i}')
        allowed = role.get('fact_ids')
        ids = requested.get(str(i), allowed)
        if not isinstance(allowed, list) or not isinstance(ids, list) or not ids:
            raise ValueError(f'role {i}: nonempty fact IDs required')
        if len(ids) != len(set(ids)) or any(fid not in allowed or fid not in by_id for fid in ids):
            raise ValueError(f'role {i}: selected fact is duplicate, unknown or outside role')
        # Selection prioritizes relevant bullets, never silently deletes source facts.
        rendered_roles.append((role, [by_id[fid] for fid in ids + [fid for fid in allowed if fid not in ids]]))
    education = profile.get('education', [])
    skills = profile.get('skills', [])
    if not isinstance(education, list) or not isinstance(skills, list) or any(not isinstance(x, str) for x in education + skills):
        raise ValueError('education and skills must be string lists')
    if not isinstance(selection.get('include_cover_letter', False), bool):
        raise ValueError('include_cover_letter must be boolean')
    cover_ids = selection.get('cover_fact_ids', [])
    highlights = profile.get('highlights', [])
    cover_sources = {**by_id}
    for item in highlights:
        required(item, ['id', 'text'], 'highlight')
        if item['id'] in cover_sources:
            raise ValueError('duplicate highlight/fact id: ' + item['id'])
        cover_sources[item['id']] = item['text']
    if (not isinstance(cover_ids, list) or len(cover_ids) > 4 or
            len(cover_ids) != len(set(cover_ids)) or
            any(not isinstance(fid, str) or fid not in cover_sources for fid in cover_ids)):
        raise ValueError('selection: invalid cover fact IDs')
    return contact, rendered_roles, education, skills


def docx_setup(doc):
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.61)
    section.left_margin = section.right_margin = Inches(.7)
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(45, 55, 65)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color, before, after in [
        ('Title', 19, RGBColor(25, 51, 72), 0, 3),
        ('Heading 1', 10.2, RGBColor(20, 97, 112), 12, 5),
        ('Heading 2', 9.7, RGBColor(25, 51, 72), 8, 3),
    ]:
        style = doc.styles[name]
        style.font.name = 'Aptos'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    bullets = doc.styles['List Bullet']
    bullets.font.name = 'Aptos'
    bullets.font.size = Pt(9.3)
    bullets.font.color.rgb = RGBColor(45, 55, 65)
    bullets.paragraph_format.space_after = Pt(3)
    bullets.paragraph_format.left_indent = Inches(.24)
    bullets.paragraph_format.first_line_indent = Inches(-.13)


def cv_blocks(profile, contact, roles, education, skills):
    blocks = [('name', profile['name']),
              ('headline', profile['headline']),
              ('contact', ' | '.join(str(contact[k]) for k in ('email', 'phone', 'location', 'linkedin', 'github') if contact.get(k)))]
    if profile.get('profile_summary'):
        blocks += [('heading', 'Profile'), ('body', profile['profile_summary'])]
    if profile.get('highlights'):
        blocks.append(('heading', 'Selected AI & Agentic Systems Work'))
        blocks.extend(('bullet', item['text']) for item in profile['highlights'])
    if profile.get('core_expertise'):
        blocks += [('heading', 'Core Expertise'), ('body', ' • '.join(profile['core_expertise']))]
    if profile.get('technology'):
        blocks.append(('heading', 'Technology'))
        blocks.extend(('body', x) for x in profile['technology'])
    blocks.append(('heading', 'Professional Experience'))
    for role, bullets in roles:
        blocks.append(('role', f"{role['title']} - {role['company']} ({role['dates']})"))
        if role.get('narrative'):
            blocks.append(('body', role['narrative']))
        blocks.extend(('bullet', text) for text in bullets)
    if profile.get('earlier_experience'):
        blocks.append(('heading', 'Earlier Experience'))
        blocks.extend(('body', x) for x in profile['earlier_experience'])
    if skills:
        blocks += [('heading', 'Skills'), ('body', ' • '.join(skills))]
    if education:
        blocks.append(('heading', 'Education'))
        blocks.extend(('body', x) for x in education)
    return blocks

def cv_docx(path, blocks):
    d = Document()
    docx_setup(d)
    for kind, text in blocks:
        if kind == 'name':
            d.add_heading(text, 0)
        elif kind in ('heading', 'role'):
            d.add_heading(text.upper() if kind == 'heading' else text,
                          level=1 if kind == 'heading' else 2)
        elif kind == 'bullet':
            d.add_paragraph(text, style='List Bullet')
        else:
            p = d.add_paragraph(text)
            if kind == 'headline':
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(20, 97, 112)
                p.paragraph_format.space_after = Pt(4)
            elif kind == 'contact':
                p.runs[0].font.size = Pt(8.2)
                p.runs[0].font.color.rgb = RGBColor(82, 94, 105)
                p.paragraph_format.space_after = Pt(9)
    d.save(path)


def letter_lines(profile, job, roles, selection):
    sources = {item['id']: item['text'] for item in profile.get('facts', []) + profile.get('highlights', [])}
    ids = selection.get('cover_fact_ids', [])
    evidence = [sources[fid] for fid in ids] if ids else [text for _, bullets in roles for text in bullets[:1]][:3]
    lines = [profile['name'], f"Application: {job['title']} — {job['company']}",
             'Dear Hiring Team,',
             f"I am applying for the {job['title']} position at {job['company']}."]
    if evidence:
        lines.append('Relevant work from my CV includes:')
        lines.extend('- ' + text for text in evidence)
    lines.extend(['I would welcome a conversation about the role.', 'Sincerely,', profile['name']])
    return lines


def letter_docx(path, lines):
    d = Document()
    docx_setup(d)
    for line in lines:
        d.add_paragraph(line)
    d.save(path)


def pdf_style():
    regular = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    bold = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont('ResumeRegular', str(regular)))
        pdfmetrics.registerFont(TTFont('ResumeBold', str(bold)))
        body_font, heading_font = 'ResumeRegular', 'ResumeBold'
    else:
        body_font, heading_font = 'Helvetica', 'Helvetica-Bold'
    ss = getSampleStyleSheet()
    navy = colors.HexColor('#193348')
    teal = colors.HexColor('#146170')
    muted = colors.HexColor('#526471')
    ss.add(ParagraphStyle(name='NameCustom', parent=ss['Title'], fontName=heading_font,
                          fontSize=19, leading=23, textColor=navy, spaceAfter=3, alignment=0))
    ss.add(ParagraphStyle(name='BodyCustom', parent=ss['Normal'], fontName=body_font,
                          fontSize=9.1, leading=12.8, textColor=navy, spaceAfter=4))
    ss.add(ParagraphStyle(name='HeadlineCustom', parent=ss['BodyCustom'], fontName=heading_font,
                          fontSize=10.3, leading=14, textColor=teal, spaceAfter=4))
    ss.add(ParagraphStyle(name='HeadingCustom', parent=ss['BodyCustom'], fontName=heading_font,
                          fontSize=10, leading=13, textColor=teal, spaceBefore=11,
                          spaceAfter=6, borderBottomWidth=.7, borderBottomColor=colors.HexColor('#C8D9DC'),
                          borderPadding=3, keepWithNext=1))
    ss.add(ParagraphStyle(name='RoleCustom', parent=ss['BodyCustom'], fontName=heading_font,
                          fontSize=9.5, leading=13, textColor=navy, spaceBefore=7, spaceAfter=3))
    ss.add(ParagraphStyle(name='BulletCustom', parent=ss['BodyCustom'], leftIndent=13,
                          firstLineIndent=-9, spaceAfter=3))
    ss.add(ParagraphStyle(name='ContactCustom', parent=ss['BodyCustom'], fontSize=8.1,
                          leading=11, textColor=muted, spaceAfter=6))
    return ss


def para(text, style):
    return Paragraph(escape(text), style)


def cv_pdf_para(kind, text, styles):
    style_names = {'name': 'NameCustom', 'headline': 'HeadlineCustom',
                   'contact': 'ContactCustom', 'heading': 'HeadingCustom',
                   'role': 'RoleCustom', 'bullet': 'BulletCustom'}
    return para(text.upper() if kind == 'heading' else text,
                styles[style_names.get(kind, 'BodyCustom')])


def build_pdf(path, blocks):
    ss = pdf_style()
    story = []
    i = 0
    while i < len(blocks):
        kind, text = blocks[i]
        if kind == 'heading' and i + 1 < len(blocks) and blocks[i + 1][0] == 'role':
            group = [cv_pdf_para(kind, text, ss)]
            i += 1
            first = True
            while i < len(blocks) and (first or blocks[i][0] not in ('role', 'heading')):
                group.append(cv_pdf_para(*blocks[i], ss))
                i += 1
                first = False
            story.append(KeepTogether(group))
            continue
        if kind == 'role':
            group = []
            first = True
            while i < len(blocks) and (first or blocks[i][0] not in ('role', 'heading')):
                group.append(cv_pdf_para(*blocks[i], ss))
                i += 1
                first = False
            story.append(KeepTogether(group))
            continue
        story.append(cv_pdf_para(kind, text, ss))
        i += 1
    SimpleDocTemplate(str(path), pagesize=(595.28, 841.89), leftMargin=48, rightMargin=48,
                      topMargin=42, bottomMargin=42, title=path.stem).build(story)


def pdf_text(path):
    return '\n'.join(page.extract_text() or '' for page in PdfReader(str(path)).pages)


def docx_text(path):
    return '\n'.join(p.text for p in Document(path).paragraphs)


def verify(path_docx, path_pdf, fragments):
    word = docx_text(path_docx)
    pdf = pdf_text(path_pdf)
    for fragment in fragments:
        # Normalize whitespace: reportlab and Word may wrap differently.
        norm = lambda x: re.sub(r'\s+', ' ', x.replace('•', ' | ')).strip().casefold()
        if norm(fragment) not in norm(word) or norm(fragment) not in norm(pdf):
            raise ValueError(f'round-trip missing text: {fragment[:80]}')
    if not pdf.strip():
        raise ValueError('empty PDF extraction')


def build(profile, job, selection, out):
    contact, roles, education, skills = assemble(profile, job, selection)
    out.mkdir(parents=True, exist_ok=True, mode=0o700)
    prefix = re.sub(r'[^A-Za-z0-9_-]+', '_', job['id']).strip('_')
    if not prefix:
        raise ValueError('job ID has no safe filename characters')
    cv_word, cv_pdf = out / f'{prefix}_CV.docx', out / f'{prefix}_CV.pdf'
    outputs = [cv_word, cv_pdf]
    if selection.get('include_cover_letter', False):
        outputs.extend((out / f'{prefix}_cover.docx', out / f'{prefix}_cover.pdf'))
    if any(path.exists() for path in outputs):
        raise FileExistsError('application artifacts already exist; choose a new revision directory')
    blocks = cv_blocks(profile, contact, roles, education, skills)
    cv_docx(cv_word, blocks)
    pdf_blocks = []
    for kind, text in blocks:
        text = text.replace(' • ', ' | ')
        if kind == 'bullet':
            text = '- ' + text
        pdf_blocks.append((kind, text))
    build_pdf(cv_pdf, pdf_blocks)
    fragments = [text for kind, text in blocks if kind not in ('contact', 'role')]
    fragments += [str(contact[k]) for k in ('email', 'phone', 'location', 'linkedin', 'github') if contact.get(k)]
    fragments += [value for role, _ in roles for value in (role['title'], role['company'], role['dates'])]
    verify(cv_word, cv_pdf, fragments)
    paths = {'cv_docx': str(cv_word), 'cv_pdf': str(cv_pdf)}
    if selection.get('include_cover_letter', False):
        lines = letter_lines(profile, job, roles, selection)
        letter_word, letter_pdf = out / f'{prefix}_cover.docx', out / f'{prefix}_cover.pdf'
        letter_docx(letter_word, lines)
        build_pdf(letter_pdf, [('name', lines[0])] + [('body', x) for x in lines[1:]])
        verify(letter_word, letter_pdf, lines)
        paths.update(cover_docx=str(letter_word), cover_pdf=str(letter_pdf))
    for path in paths.values():
        Path(path).chmod(0o600)
    return paths


def main():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest='command', required=True)
    importer = sub.add_parser('import-source')
    importer.add_argument('--source', required=True)
    importer.add_argument('--out', required=True)
    b = sub.add_parser('build')
    for arg in ('profile', 'job', 'selection', 'out-dir'):
        b.add_argument('--' + arg, required=True)
    args = p.parse_args()
    if args.command == 'import-source':
        result = import_source(args.source)
        target = Path(args.out)
        target.parent.mkdir(parents=True, exist_ok=True)
        with os.fdopen(os.open(target, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600), 'w', encoding='utf-8') as stream:
            json.dump(result, stream, ensure_ascii=False, indent=2)
            stream.write('\n')
        target.chmod(0o600)
        print(json.dumps({'profile': str(target.resolve()), 'roles': len(result['roles']),
                          'highlights': len(result['highlights']), 'earlier_experience': len(result['earlier_experience'])}))
    else:
        print(json.dumps(build(load(args.profile), load(args.job), load(args.selection), Path(args.out_dir)), indent=2))


if __name__ == '__main__':
    main()
