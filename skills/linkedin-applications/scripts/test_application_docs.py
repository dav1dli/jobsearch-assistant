"""Run with the skill venv: python scripts/test_application_docs.py."""
import copy
import importlib.util
import tempfile
from pathlib import Path

spec = importlib.util.spec_from_file_location('application_docs', Path(__file__).with_name('application_docs.py'))
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)

profile = {
    'name': 'Test Candidate', 'headline': 'Technology leader',
    'contact': {'email': 'example@example.org', 'phone': '+1 555 0100', 'location': 'Test City'},
    'facts': [{'id': 'a', 'text': 'Led a platform migration.'}, {'id': 'b', 'text': 'Mentored engineers.'}],
    'roles': [{'title': 'Architect', 'company': 'Example Co', 'dates': '2020–2024', 'fact_ids': ['a', 'b']}],
    'education': ['Example University'], 'skills': ['Architecture', 'Mentoring']
}
job = {'id': 'test-job', 'title': 'CTO', 'company': 'Example Corp',
       'url': 'https://example.org/jobs/1', 'location': 'Test City', 'work_model': 'hybrid'}
selection = {'role_fact_ids': {'0': ['a']}, 'include_cover_letter': True}
with tempfile.TemporaryDirectory() as directory:
    first_revision = Path(directory) / 'revision-01'
    files = module.build(profile, job, selection, first_revision)
    assert len(files) == 4 and all(Path(f).exists() for f in files.values())
    assert 'Mentored engineers' in module.docx_text(files['cv_docx'])
    assert 'Led a platform migration' in module.pdf_text(files['cv_pdf'])
    tailored = module.build(profile, job, {'include_cover_letter': True,
                                          'cover_fact_ids': ['b']}, Path(directory) / 'revision-02')
    assert 'Mentored engineers.' in module.pdf_text(tailored['cover_pdf'])
    assert 'Led a platform migration.' not in module.pdf_text(tailored['cover_pdf'])
    try:
        module.build(profile, job, {'include_cover_letter': True,
                                    'cover_fact_ids': ['invented']}, Path(directory) / 'invalid-cover')
    except ValueError:
        pass
    else:
        raise AssertionError('unverified cover claim was accepted')
    bad = copy.deepcopy(selection)
    bad['role_fact_ids']['0'] = ['invalid']
    try:
        module.build(profile, job, bad, Path(directory) / 'invalid-role')
    except ValueError:
        pass
    else:
        raise AssertionError('unverified fact was accepted')
    no_letter = module.build(profile, job, {'role_fact_ids': {'0': ['a']}}, Path(directory) / 'revision-no-letter')
    assert set(no_letter) == {'cv_docx', 'cv_pdf'}
    try:
        module.build(profile, job, selection, first_revision)
    except FileExistsError:
        pass
    else:
        raise AssertionError('existing artifacts were overwritten')
print('PASS: 4 artifacts, text round-trip, fact guard, optional letter')

# A source-faithful profile retains every section and role narrative even when
# selected IDs are used to prioritize a particular role's bullets.
rich = copy.deepcopy(profile)
rich['profile_summary'] = 'Built products from inception.'
rich['highlights'] = [{'id': 'h1', 'text': 'Agent gateway with durable tasks.', 'source_paragraph': 6}]
rich['core_expertise'] = ['Agent systems', 'Enterprise architecture']
rich['technology'] = ['Languages: Python, Java', 'Platforms: Linux']
rich['roles'][0]['narrative'] = 'Architected a production platform.'
rich['earlier_experience'] = ['CTO, Earlier Co | 2010–2012 — Built a product.']
rich['contact']['github'] = 'github.com/test'
with tempfile.TemporaryDirectory() as directory:
    files = module.build(rich, job, selection, Path(directory))
    assert all(Path(f).stat().st_mode & 0o077 == 0 for f in files.values())
    for key in ('cv_docx', 'cv_pdf'):
        text = module.docx_text(files[key]) if key.endswith('docx') else module.pdf_text(files[key])
        for expected in ('Built products from inception.', 'Agent gateway with durable tasks.',
                         'Enterprise architecture', 'Platforms: Linux',
                         'Architected a production platform.', 'CTO, Earlier Co',
                         'Mentored engineers.', 'github.com/test'):
            assert expected in text, (key, expected)
    no_email = copy.deepcopy(rich)
    del no_email['contact']['email']
    files = module.build(no_email, job, {}, Path(directory) / 'revision-no-email')
    assert 'example@example.org' not in module.docx_text(files['cv_docx'])
print('PASS: source sections, narrative, selection retention, absent email')

from docx import Document
with tempfile.TemporaryDirectory() as directory:
    source = Path(directory) / 'source.docx'
    d = Document()
    for text, style in [('Test Candidate', 'Title'), ('Technology leader', 'Normal'),
                        ('Test City | +1 555 0100 | github.com/test', 'Normal'),
                        ('PROFILE', 'Heading 1'), ('Built products.', 'Normal'),
                        ('SELECTED AI & AGENTIC SYSTEMS WORK', 'Heading 1'),
                        ('• Agent gateway.', 'Normal'), ('CORE EXPERTISE', 'Heading 1'),
                        ('Agents • Systems', 'Normal'), ('TECHNOLOGY', 'Heading 1'),
                        ('Languages: Python', 'Normal'), ('PROFESSIONAL EXPERIENCE', 'Heading 1'),
                        ('Architect — Example Co', 'Heading 2'), ('2020 – 2024', 'Normal'),
                        ('Led design.', 'Normal'), ('• Built services.', 'Normal'),
                        ('EARLIER EXPERIENCE', 'Heading 1'),
                        ('CTO, Old Co | 2010 – 2012 — Built product.', 'Normal'),
                        ('EDUCATION', 'Heading 1'), ('B.Sc., Example U', 'Normal')]:
        d.add_paragraph(text, style)
    d.save(source)
    imported = module.import_source(source)
    assert imported['source']['sha256'] and imported['source']['path'] == str(source.resolve())
    assert imported['roles'][0]['narrative'] == 'Led design.'
    assert d.paragraphs[imported['facts'][0]['source_paragraph']].text == '• Built services.'
    assert imported['highlights'][0]['text'] == 'Agent gateway.'
    assert imported['earlier_experience'][0].startswith('CTO, Old Co')
    assert 'email' not in imported['contact']
print('PASS: source import with paragraph provenance and no inferred email')

markdown = '''# Test Candidate
Technology leader
example@example.org | +1 555 0100 | Test City

## Summary
Built products from inception.

## Core Expertise
Architecture • Mentoring

## Professional Experience
### Architect — Example Co
2020–2024
Led a platform migration.
- Mentored engineers.

## Education
Example University
'''
with tempfile.TemporaryDirectory() as directory:
    source = Path(directory) / 'source.md'
    source.write_text(markdown, encoding='utf-8')
    imported = module.import_source(source)
    assert imported['source']['format'] == 'markdown'
    assert imported['name'] == 'Test Candidate'
    assert imported['roles'][0]['company'] == 'Example Co'
    assert imported['facts'][0]['source_location']['kind'] == 'line'

from reportlab.pdfgen import canvas
with tempfile.TemporaryDirectory() as directory:
    source = Path(directory) / 'source.pdf'
    pdf = canvas.Canvas(str(source))
    for number, line in enumerate([
        'Test Candidate', 'Technology leader', 'example@example.org | +1 555 0100 | Test City',
        'PROFILE', 'Built products from inception.', 'CORE EXPERTISE', 'Architecture - Mentoring',
        'PROFESSIONAL EXPERIENCE', 'Architect - Example Co', '2020-2024',
        'Led a platform migration.', '- Mentored engineers.', 'EDUCATION', 'Example University',
    ]):
        pdf.drawString(50, 800 - number * 35, line)
    pdf.save()
    imported = module.import_source(source)
    assert imported['source']['format'] == 'pdf'
    assert imported['roles'][0]['company'] == 'Example Co'
    assert imported['facts'][0]['source_location']['kind'] == 'page_line'

    scanned = Path(directory) / 'scanned.pdf'
    blank = canvas.Canvas(str(scanned))
    blank.showPage()
    blank.save()
    try:
        module.import_source(scanned)
    except ValueError as error:
        assert 'no selectable text' in str(error)
    else:
        raise AssertionError('image-only PDF was silently accepted')
print('PASS: Markdown and selectable PDF import with source locators; scan-only PDF fails clearly')
